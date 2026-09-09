#!/usr/bin/env python3
"""
Video Chat Application PowerPoint Generator - Simplified Version
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

def create_video_chat_presentation():
    """Create the complete Video Chat Application presentation"""
    print("🎬 Creating Video Chat Application Presentation...")
    
    # Create presentation
    prs = Presentation()
    
    # Color scheme
    deep_blue = RGBColor(30, 58, 138)
    bright_orange = RGBColor(249, 115, 22)
    clean_white = RGBColor(255, 255, 255)
    dark_gray = RGBColor(55, 65, 81)
    
    # Slide 1: Title Slide
    print("📝 Creating slide 1: Title...")
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "The Future of Learning: When Videos Talk Back"
    title_frame = title.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(48)
    title_para.font.color.rgb = deep_blue
    title_para.font.bold = True
    
    subtitle = slide.shapes.placeholders[1]
    subtitle.text = "Transforming One-Way Video Learning into Interactive Conversations\n\nJune 20, 2025"
    subtitle_frame = subtitle.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(24)
    subtitle_para.font.color.rgb = dark_gray
    
    # Slide 2: Opening Hook
    print("📝 Creating slide 2: Hook...")
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "Did You Know?"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1]
    content.text = "🙋 Average person watches 17 hours of video content weekly\n\n😕 Most learning videos are one-way communication\n\n🤔 What if your videos could talk back?"
    content.text_frame.paragraphs[0].font.size = Pt(28)
    content.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Slide 3: The Problem
    print("📝 Creating slide 3: Problem...")
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "The Learning Problem"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1]
    content.text = "📺 Learning = One-Way Street\n\n❓ Questions? → Rewind → Research → Lost Focus\n\n⏰ Time wasted searching for specific information\n\n😤 Frustration when videos can't answer back"
    content.text_frame.paragraphs[0].font.size = Pt(28)
    content.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Slide 4: The Solution
    print("📝 Creating slide 4: Solution...")
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "Introducing Video Chat App"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1]
    content.text = "🎬 Any YouTube Video → Intelligent Companion\n\n🧠 Powered by AI & Built for Learners\n\n💬 Chat about anything in the video\n\n⚡ Science fiction that runs on your laptop today"
    content.text_frame.paragraphs[0].font.size = Pt(28)
    content.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Slide 5: How It Works
    print("📝 Creating slide 5: How it works...")
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "How It Works"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1]
    content.text = "1️⃣ Paste YouTube URL\n\n2️⃣ AI Extracts & Analyzes Content\n\n3️⃣ Chat About Anything in the Video\n\n✨ Get instant, contextual responses!"
    content.text_frame.paragraphs[0].font.size = Pt(32)
    content.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Slide 6: Key Features
    print("📝 Creating slide 6: Features...")
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "Key Features"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1]
    content.text = "🎯 Contextual AI Chat - Answers based on video content\n\n⏱️ Auto-Generated Timestamps - AI creates chapter markers\n\n🔍 Instant Content Search - Find anything in seconds\n\n📺 Universal Compatibility - Works with ANY YouTube video"
    content.text_frame.paragraphs[0].font.size = Pt(26)
    content.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Slide 7: What Makes It Special
    print("📝 Creating slide 7: Special features...")
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "What Makes It Special"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1]
    content.text = "🌍 Universal Compatibility - No special setup needed\n\n🧠 Context-Aware Intelligence - Real video understanding\n\n❤️ Built by Learners, for Learners - Solving real problems"
    content.text_frame.paragraphs[0].font.size = Pt(30)
    content.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Slide 8: The Impact
    print("📝 Creating slide 8: Impact...")
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "The Impact We're Creating"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    content = slide.shapes.placeholders[1]
    content.text = "🔄 Making Every Video Interactive\n\n💡 Turning Confusion into Understanding\n\n🌉 Breaking Learning Barriers\n\n🚀 Transforming Education Technology"
    content.text_frame.paragraphs[0].font.size = Pt(30)
    content.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Slide 9: Closing Quote
    print("📝 Creating slide 9: Closing...")
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Set background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = deep_blue
    
    # Add quote
    quote_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(11.33), Inches(3))
    quote_frame = quote_box.text_frame
    quote_frame.word_wrap = True
    
    # First line
    para1 = quote_frame.paragraphs[0]
    para1.text = "Learning shouldn't be a monologue."
    para1.font.size = Pt(36)
    para1.font.color.rgb = clean_white
    para1.font.bold = True
    para1.alignment = PP_ALIGN.CENTER
    
    # Second line
    para2 = quote_frame.add_paragraph()
    para2.text = "It should be a conversation."
    para2.font.size = Pt(36)
    para2.font.color.rgb = bright_orange
    para2.font.bold = True
    para2.alignment = PP_ALIGN.CENTER
    
    # Third line
    para3 = quote_frame.add_paragraph()
    para3.text = "And now, finally, it can be."
    para3.font.size = Pt(36)
    para3.font.color.rgb = clean_white
    para3.font.bold = True
    para3.alignment = PP_ALIGN.CENTER
    
    # Slide 10: Thank You
    print("📝 Creating slide 10: Thank you...")
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    title = slide.shapes.title
    title.text = "Thank You"
    title.text_frame.paragraphs[0].font.size = Pt(60)
    title.text_frame.paragraphs[0].font.color.rgb = deep_blue
    title.text_frame.paragraphs[0].font.bold = True
    
    subtitle = slide.shapes.placeholders[1]
    subtitle.text = "Questions?\n\nDemo available at: [your-demo-link]\nContact: [your-email]"
    subtitle.text_frame.paragraphs[0].font.size = Pt(28)
    subtitle.text_frame.paragraphs[0].font.color.rgb = dark_gray
    
    # Save presentation
    output_path = "Video_Chat_Application_Presentation.pptx"
    print(f"💾 Saving presentation as {output_path}...")
    prs.save(output_path)
    
    print("✅ Presentation created successfully!")
    print(f"📁 File: {output_path}")
    print(f"📊 Total slides: {len(prs.slides)}")
    
    return output_path

if __name__ == "__main__":
    try:
        result = create_video_chat_presentation()
        print(f"\n🎉 SUCCESS! Your presentation is ready: {result}")
        print("\n📋 Next steps:")
        print("1. Open the PowerPoint file")
        print("2. Add your personal details")
        print("3. Practice with the delivery guide")
        print("4. Add demo videos if desired")
    except Exception as e:
        print(f"❌ Error: {str(e)}")
        import traceback
        traceback.print_exc()
